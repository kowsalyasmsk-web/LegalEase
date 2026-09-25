from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from PIL import Image as PILImage


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",
        "\u2019": "'",

        "\u201c": '"',
        "\u201d": '"',

        "\u2013": "-",
        "\u2014": "-",

        "\u00a0": " ",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    return text.strip()


def safe_filename(value: str) -> str:

    value = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        value.strip()
    )

    return (
        value.strip("._")
        or "legal_document"
    )


def filename_for(
    document_type: str,
    extension: str
) -> str:

    return (
        f"{safe_filename(document_type)}"
        f"_LegalEase.{extension}"
    )


def format_txt(text: str) -> bytes:

    clean_text = sanitize_text(
        text
    )

    return clean_text.encode(
        "utf-8"
    )


def format_docx(
    text: str,
    document_type: str,
    logo_bytes: bytes | None = None
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal_style = document.styles["Normal"]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(11)

    # Logo

    if logo_bytes:

        try:

            document.add_picture(
                BytesIO(logo_bytes),
                width=Inches(1.5)
            )

            document.paragraphs[-1].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

        except Exception:
            pass

    # Title

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        document_type.upper()
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(16)

    # Content

    for raw_line in sanitize_text(
        text
    ).splitlines():

        line = raw_line.strip()

        if not line:

            document.add_paragraph("")

            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = (
            Pt(6)
        )

        is_heading = (

            len(line) < 100

            and (

                line.isupper()

                or re.match(
                    r"^\d+[\.)]\s+",
                    line
                )

                or line.lower().endswith(":")
            )
        )

        run = paragraph.add_run(
            line
        )

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(11)

        run.bold = is_heading

    # Footer

    footer = (
        document.sections[0]
        .footer
        .paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase — AI-generated draft. "
        "Review before signing."
    )

    footer_run.font.name = (
        "Times New Roman"
    )

    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(
        output
    )

    return output.getvalue()


def pdf_header_footer(
    canvas,
    doc,
    logo_bytes
):

    canvas.saveState()

    width, height = A4

    if logo_bytes:

        try:

            image = PILImage.open(
                BytesIO(logo_bytes)
            )

            image.thumbnail(
                (120, 60)
            )

            temp = BytesIO()

            image.save(
                temp,
                format="PNG"
            )

            temp.seek(0)

            canvas.drawImage(

                temp,

                width / 2 - 45,

                height - 65,

                width=90,

                height=45,

                preserveAspectRatio=True,

                mask="auto",
            )

        except Exception:
            pass

    canvas.setFont(
        "Helvetica",
        8
    )

    canvas.drawCentredString(

        width / 2,

        25,

        "LegalEase — AI-generated draft. "
        "Review before signing.",
    )

    canvas.restoreState()


def format_pdf(
    text: str,
    document_type: str,
    logo_bytes: bytes | None = None
) -> bytes:

    output = BytesIO()

    document = SimpleDocTemplate(

        output,

        pagesize=A4,

        rightMargin=50,

        leftMargin=50,

        topMargin=60,

        bottomMargin=45,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(

        "LegalTitle",

        parent=styles["Title"],

        alignment=TA_CENTER,

        fontName="Helvetica-Bold",

        fontSize=16,

        spaceAfter=18,
    )

    heading_style = ParagraphStyle(

        "LegalHeading",

        parent=styles["Heading2"],

        fontName="Helvetica-Bold",

        fontSize=11,

        spaceBefore=8,

        spaceAfter=5,
    )

    body_style = ParagraphStyle(

        "LegalBody",

        parent=styles["BodyText"],

        fontName="Times-Roman",

        fontSize=10.5,

        leading=15,

        spaceAfter=7,
    )

    story = [

        Paragraph(
            document_type.upper(),
            title_style
        )
    ]

    for raw_line in sanitize_text(
        text
    ).splitlines():

        line = raw_line.strip()

        if not line:

            story.append(
                Spacer(1, 5)
            )

            continue

        escaped = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        is_heading = (

            len(line) < 100

            and (

                line.isupper()

                or re.match(
                    r"^\d+[\.)]\s+",
                    line
                )

                or line.lower().endswith(":")
            )
        )

        style = (
            heading_style
            if is_heading
            else body_style
        )

        story.append(
            Paragraph(
                escaped,
                style
            )
        )

    document.build(

        story,

        onFirstPage=lambda canvas, doc:
            pdf_header_footer(
                canvas,
                doc,
                logo_bytes
            ),

        onLaterPages=lambda canvas, doc:
            pdf_header_footer(
                canvas,
                doc,
                logo_bytes
            ),
    )

    return output.getvalue()